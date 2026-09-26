"""
Test Suite for MAANAKNETRA Document Ingestion Layer.
Covers:
1. TXT extraction
2. PDF extraction
3. DOCX extraction
4. Empty document handling
5. Scanned / near-empty PDF OCR fallback decision
6. Page boundary preservation
Integration test with synthetic sample tender.
"""

import os
import sys
import unittest
import tempfile
import pymupdf
from docx import Document

# Ensure root workspace is on python sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from backend.ingestion import (
    ingest_document,
    should_trigger_ocr,
    ExtractionMethod,
    FileType
)


class TestDocumentIngestion(unittest.TestCase):

    def setUp(self):
        self.sample_txt = os.path.join("demo", "sample_tender.txt")
        self.sample_pdf = os.path.join("demo", "sample_tender.pdf")
        self.assertTrue(os.path.exists(self.sample_txt), "sample_tender.txt must exist")
        self.assertTrue(os.path.exists(self.sample_pdf), "sample_tender.pdf must exist")

    def test_1_txt_extraction(self):
        """Test TXT parsing, whitespace normalization, and clause detection."""
        result = ingest_document(self.sample_txt)
        self.assertEqual(result.file_type, "TXT")
        self.assertGreater(result.page_count, 0)
        self.assertIn("RWSS/ZONE-IV/PUMP-PROC/2026/088", result.extracted_text)
        self.assertIn("IS 14220:1994", result.extracted_text)

        # Check that blocks and clauses were recognized
        first_page = result.pages[0]
        self.assertEqual(first_page.extraction_method, ExtractionMethod.TEXT)
        self.assertGreater(len(first_page.blocks), 0)

        # Verify a specific clause was extracted
        clauses = [b.clause_number for b in first_page.blocks if b.clause_number]
        self.assertTrue(any("1.1" in c or "SECTION" in c for c in clauses))

    def test_2_pdf_extraction(self):
        """Test PDF parsing via PyMuPDF on the synthetic sample tender."""
        result = ingest_document(self.sample_pdf)
        self.assertEqual(result.file_type, "PDF")
        self.assertEqual(result.page_count, 3)
        self.assertIn("SYNTHETIC DEMONSTRATION DATA", result.extracted_text)
        self.assertIn("IS 14220:1994", result.extracted_text)

        # Verify page-level details
        for p in result.pages:
            self.assertEqual(p.extraction_method, ExtractionMethod.TEXT)
            self.assertGreater(p.character_count, 100)
            self.assertGreater(len(p.blocks), 0)

    def test_3_docx_extraction(self):
        """Test DOCX parsing with paragraphs, page breaks, and tables."""
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp_docx_path = tmp.name

        try:
            doc = Document()
            doc.add_heading("STATE WATER SUPPLY TENDER", level=1)
            doc.add_paragraph("Clause 1.1: The contractor shall supply 50 submersible pumpsets.")
            doc.add_paragraph("Clause 1.2: Governing specification shall be IS 14220:2018.")
            doc.add_page_break()
            doc.add_paragraph("Clause 2.1: Tested in accordance with IS 11346.")

            # Add a small parameter table
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Parameter"
            table.cell(0, 1).text = "Required Value"
            table.cell(1, 0).text = "Operating Voltage"
            table.cell(1, 1).text = "415 V"

            doc.save(tmp_docx_path)

            result = ingest_document(tmp_docx_path)
            self.assertEqual(result.file_type, "DOCX")
            self.assertEqual(result.page_count, 2)
            self.assertIn("STATE WATER SUPPLY TENDER", result.extracted_text)
            self.assertIn("IS 14220:2018", result.extracted_text)
            self.assertIn("Operating Voltage | 415 V", result.extracted_text)

        finally:
            if os.path.exists(tmp_docx_path):
                os.remove(tmp_docx_path)

    def test_4_empty_document(self):
        """Test graceful handling of a 0-byte or completely empty document."""
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as tmp:
            tmp_empty_path = tmp.name

        try:
            result = ingest_document(tmp_empty_path)
            self.assertEqual(result.character_count, 0)
            self.assertEqual(result.extracted_text, "")
            self.assertEqual(result.metadata.get("status"), "EMPTY_DOCUMENT")
        finally:
            if os.path.exists(tmp_empty_path):
                os.remove(tmp_empty_path)

    def test_5_scanned_near_empty_pdf_ocr_decision(self):
        """Test that OCR fallback decision triggers when PDF page has little/no extractable text."""
        # 1. Direct unit test of trigger threshold
        self.assertTrue(should_trigger_ocr(""))
        self.assertTrue(should_trigger_ocr("     \n\n  "))
        self.assertTrue(should_trigger_ocr("Scanned Doc"))  # 11 chars < 40 chars
        self.assertFalse(should_trigger_ocr("This is a full page containing well over forty characters of digital text."))

        # 2. Integration test: Generate a synthetic PDF page with near-empty text
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            tmp_scanned_pdf = tmp.name

        try:
            doc = pymupdf.open()
            page = doc.new_page(width=595, height=842)
            # Add minimal text (< 20 chars) simulating an image stamp or blank scanned page
            page.insert_textbox(pymupdf.Rect(50, 50, 200, 100), "Stamp 12", fontsize=10)
            doc.save(tmp_scanned_pdf)
            doc.close()

            result = ingest_document(tmp_scanned_pdf, ocr_threshold=40)
            # Should detect low text count and evaluate OCR without crashing
            self.assertEqual(result.page_count, 1)
            self.assertIsNotNone(result.pages[0])

        finally:
            if os.path.exists(tmp_scanned_pdf):
                os.remove(tmp_scanned_pdf)

    def test_6_page_boundary_preservation(self):
        """Test that individual page boundaries and numbering are strictly preserved."""
        result = ingest_document(self.sample_pdf)
        self.assertEqual(len(result.pages), 3)

        # Page numbers must be consecutive 1-indexed
        for idx, page in enumerate(result.pages, start=1):
            self.assertEqual(page.page_number, idx)
            self.assertGreater(page.character_count, 0)
            # Verify blocks all map back to this page
            for block in page.blocks:
                self.assertEqual(block.page_number, idx)

        # Confirm different sections landed on expected pages
        self.assertIn("SECTION 1", result.pages[0].text)
        self.assertIn("SECTION 4", result.pages[1].text)
        self.assertIn("SECTION 6", result.pages[2].text)


if __name__ == "__main__":
    unittest.main()
