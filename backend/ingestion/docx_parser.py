"""
DOCX Document Parser for MAANAKNETRA.
Parses Word documents (.docx), extracts paragraphs, tables, clause numbers, and handles page breaks.
"""

import os
import re
from typing import List, Tuple
from .models import PageObject, TextBlock, ExtractionMethod

CLAUSE_REGEX = re.compile(r"^(\d+(\.\d+)*|[A-Z](\.\d+)*|Clause\s+\d+(\.\d+)*|SECTION\s+\d+)", re.IGNORECASE)


def parse_docx_file(file_path: str) -> List[PageObject]:
    """
    Parses a DOCX document into structured PageObjects with TextBlocks.
    Handles explicit page breaks and tables.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    try:
        from docx import Document
    except ImportError:
        raise RuntimeError("python-docx is not installed. Install via pip install python-docx")

    try:
        doc = Document(file_path)
    except Exception as e:
        raise ValueError(f"Corrupted or invalid DOCX document {file_path}: {e}")

    # Segment content into pages by inspecting paragraph text and explicit page breaks
    pages_raw_elements: List[List[str]] = [[]]

    for element in doc.paragraphs:
        # Check if paragraph has an explicit page break
        has_page_break = False
        for run in element.runs:
            if 'w:br' in run._element.xml and 'w:type="page"' in run._element.xml:
                has_page_break = True
                break

        text = element.text.strip()
        if text:
            pages_raw_elements[-1].append(text)

        if has_page_break:
            pages_raw_elements.append([])

    # Also capture table contents
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                table_rows.append(" | ".join(row_cells))
        if table_rows:
            table_text = "\n".join(table_rows)
            pages_raw_elements[-1].append(f"[TABLE]\n{table_text}")

    # Remove trailing empty pages if any
    if len(pages_raw_elements) > 1 and not pages_raw_elements[-1]:
        pages_raw_elements.pop()

    pages: List[PageObject] = []

    for page_idx, elements in enumerate(pages_raw_elements, start=1):
        clean_page_text, blocks = _build_blocks_from_elements(elements, page_idx)
        pages.append(PageObject(
            page_number=page_idx,
            text=clean_page_text,
            extraction_method=ExtractionMethod.TEXT,
            character_count=len(clean_page_text),
            blocks=blocks
        ))

    # If document has no elements at all, return a single empty page
    if not pages:
        pages.append(PageObject(
            page_number=1,
            text="",
            extraction_method=ExtractionMethod.TEXT,
            character_count=0,
            blocks=[]
        ))

    return pages


def _build_blocks_from_elements(elements: List[str], page_number: int) -> Tuple[str, List[TextBlock]]:
    blocks: List[TextBlock] = []
    normalized_parts: List[str] = []
    current_char_offset = 0

    for idx, text in enumerate(elements, start=1):
        clean_text = text.strip()
        if not clean_text:
            continue

        clause_match = CLAUSE_REGEX.match(clean_text)
        clause_num = clause_match.group(1) if clause_match else None

        char_start = current_char_offset
        char_end = char_start + len(clean_text)

        blocks.append(TextBlock(
            block_id=f"p{page_number}_b{idx}",
            clause_number=clause_num,
            text=clean_text,
            page_number=page_number,
            char_start=char_start,
            char_end=char_end
        ))

        normalized_parts.append(clean_text)
        current_char_offset = char_end + 2

    full_text = "\n\n".join(normalized_parts)
    return full_text, blocks
